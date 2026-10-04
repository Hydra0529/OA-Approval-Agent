from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Cm

desk = Path(r"C:\Users\Jhy13\Desktop")
EAST_ASIA = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia"


def _font(run, size=12, bold=False, italic=False, color=None):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(EAST_ASIA, "宋体")
    if color is not None:
        run.font.color.rgb = color


def add_h(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_p(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _font(run, bold=bold)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.3


def add_bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(it, style="List Bullet")
        for run in p.runs:
            _font(run)


def add_num(doc, items):
    for it in items:
        p = doc.add_paragraph(it, style="List Number")
        for run in p.runs:
            _font(run)


def add_code(doc, text):
    for line in text.strip("\n").splitlines():
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(10)
        run._element.rPr.rFonts.set(EAST_ASIA, "宋体")
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
    doc.add_paragraph()


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                _font(run, size=11, bold=True)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    _font(run, size=11)
    doc.add_paragraph()


def main():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.2)
        s.bottom_margin = Cm(2.2)
        s.left_margin = Cm(2.5)
        s.right_margin = Cm(2.5)

    doc.add_heading("公司 OA 审批 AI Agent：多方案与首期落地计划", level=0)
    add_p(
        doc,
        "概述：面向通用公司 OA 审批场景，首期目标为交付「项目群清单 + 基于 PDF 制度学习的流程路径/时间线/BPMN/分析报告」演示系统。流程拓扑以规章制度为唯一知识源，不依赖 OA 事件日志挖掘。",
    )

    add_h(doc, "一、需求锚定")
    add_bullets(
        doc,
        [
            "业务域：通用公司 OA（报销 / 采购 / 合同 / 请假等），不是对标视频的长周期军建仿真。",
            "首期交付：一份可交互的项目群清单（多条审批案例列表：名称、类型、主责/经办、预计与实际时间、累计时长、当前节点与下一步时间线）。",
            "流程路径来源：只从 PDF 等规章制度学习（制度 → 节点/角色/条件/SLA），不用 OA 流水做路径发现。",
            "约束：自备演示数据；不涉及隐私与微调；可用云端开源/开放权重大模型（DeepSeek / Qwen 等 OpenAI 兼容 API）。",
        ],
    )
    add_p(doc, "视频里要复现的「扭转机制」抽象为四件事：")
    add_num(
        doc,
        [
            "制度 → 结构化流程（节点、网关、角色、时限）",
            "实例进入某节点后 → 按规则推下一步与预计日期",
            "多实例汇总 → 项目群清单 + 态势",
            "单实例/全流程 → BPMN 图 + 分析报告",
        ],
    )
    add_p(doc, "总体数据流：")
    add_code(
        doc,
        """制度PDF → 解析切块 → 向量/检索知识库
              ↘ LLM结构化抽取 → 流程模板库
模拟审批实例 + 流程模板 → 监控/推演引擎
  → 项目群清单 / 下一步时间线
流程模板 → BPMN渲染
检索知识库 → 解释与分析报告""",
    )

    add_h(doc, "二、四套方案对比")

    add_h(doc, "方案 A — 纯 RAG Agent（文档问答驱动）", 2)
    add_bullets(
        doc,
        [
            "做法：PDF 入库检索；用户/系统传入「当前单据摘要」，Agent 直接生成下一步、时间线、BPMN XML、报告。",
            "优点：实现最快，最像「AI 读制度」。",
            "缺点：路径与日期易幻觉；BPMN 常不合法；清单字段难稳定。",
            "适用：口头售前 demo，不推荐作首期主路径。",
        ],
    )

    add_h(doc, "方案 B — 制度抽取 → 流程模板 + 轻量推演（推荐首期）", 2)
    add_p(doc, "做法：")
    add_num(
        doc,
        [
            "PDF 切块 RAG（解释「为何走这条」）。",
            "LLM 一次性/分章抽取为可执行流程模板（JSON/YAML：节点、角色、条件、默认 SLA 天）。",
            "确定性推演引擎根据模板 + 实例字段算下一节点与日期（如「2026-01-20 提交财务初审，由财务专员审批」）。",
            "模板转 BPMN XML，前端 bpmn-js 只读展示；LLM 写分析报告（引用制度片段）。",
            "首期 UI 核心就是项目群清单表 + 行详情（时间线 / BPMN / 报告）。",
        ],
    )
    add_bullets(
        doc,
        [
            "优点：扭转可审计、可回放；与视频「规则推演」同构，但规则来自制度抽取而非手搓；符合「路径不学自 OA」。",
            "缺点：需人审一版模板金标准；制度变更要重抽。",
            "适用：首期默认选型。",
        ],
    )

    add_h(doc, "方案 C — 文档优先工作台（上传 PDF/审批单 → 清单）", 2)
    add_bullets(
        doc,
        [
            "做法：偏交互：上传制度包 → 建模板；上传「模拟审批单 PDF」→ 填实例 → 进清单。无 Webhook。",
            "优点：零对接成本，演示叙事清晰。",
            "缺点：与「实时提交监控」叙事弱；可与 B 合并（B 的交互外壳）。",
            "适用：作为方案 B 的产品形态，不单独成架构。",
        ],
    )

    add_h(doc, "方案 D — 流程挖掘优先（pm4py）", 2)
    add_bullets(
        doc,
        [
            "做法：从事件日志发现 BPMN/瓶颈。",
            "结论：首期明确不做——与「路径从 PDF 学、不从 OA 学」冲突；后期若有历史日志可作增强，不进 MVP。",
        ],
    )
    add_p(
        doc,
        "选型结论：首期落地 方案 B（内核）+ 方案 C（交互）；A 仅作辅助问答；D 搁置。",
        bold=True,
    )

    add_h(doc, "三、推荐技术栈")
    add_table(
        doc,
        ["层", "选型"],
        [
            ["后端", "Python 3.11+ / FastAPI"],
            ["Agent/编排", "LangGraph 或轻量自研 tool-calling 循环"],
            ["LLM", "DeepSeek 或 Qwen（OpenAI 兼容）；环境变量切换，不绑定一家"],
            ["PDF", "PyMuPDF + 按章节/标题切块"],
            ["检索/知识库", "首期本地索引即可；生产可换向量库"],
            ["模板/实例存储", "JSON/YAML 文件（流程模板、项目群清单）；可扩展 SQLite"],
            ["BPMN", "模板 → BPMN 2.0 XML（确定性生成）；前端 bpmn-js 只读"],
            ["前端", "React + TypeScript；清单表为主界面"],
            ["部署", "单机 docker compose 或本地 uvicorn + vite"],
        ],
    )
    add_p(
        doc,
        "硬约束：LLM 不直接改实例状态；只负责抽取模板草案、解释、写报告。推演与下一节点由模板引擎计算。",
        bold=True,
    )

    add_h(doc, "四、首期产品范围（MVP）")
    add_h(doc, "范围包含", 2)
    add_num(
        doc,
        [
            "制度知识库：导入 8–15 份自备制度 PDF（报销、采购、合同、请假等），切块入库，问答带引用。",
            "流程模板抽取：从指定制度生成/更新 YAML 模板（节点序列、条件分支、角色、SLA）；提供「人审确认」开关（演示可一键确认）。",
            "模拟项目群清单：预置 8–20 条合成审批实例（覆盖 3–5 类流程、不同金额/部门以触发分支）；表格字段含单据名称、流程类型、级别或金额档、申请人部门、当前审批人角色、预计/实际时间、累计时长、当前节点、下一步、预计到达日。",
            "行详情：时间线列表 + BPMN 图 + 简短分析报告（瓶颈节点、模板 SLA、制度依据）。",
            "模拟提交：表单或 JSON 提交一条新单据 → 立刻进清单并给出下一步时间线（证明「监控」能力，无需真 OA）。",
        ],
    )
    add_h(doc, "首期不做", 2)
    add_bullets(
        doc,
        [
            "真实 OA/Webhook 对接、回写审批",
            "从 OA 日志做流程挖掘",
            "模型微调 / 训练",
            "风险注入推演、25x 时间轴播放（可二期仿视频增强）",
        ],
    )

    add_h(doc, "五、自备数据清单")
    add_h(doc, "制度 PDF（学习源）", 2)
    add_bullets(
        doc,
        [
            "费用报销管理办法、差旅标准、采购/合同审批权限表、请假制度、印章/用印流程等",
            "要求：正文含「谁批、阈值条件、时限（工作日）」——否则无法稳定抽 SLA",
        ],
    )
    add_h(doc, "合成实例（清单用，非路径学习源）", 2)
    add_bullets(
        doc,
        [
            "每类流程 3–5 条 case：字段含 case_id / type / amount / dept / submit_at / current_node / status",
            "用模板引擎按 SLA 回填「预计完成」；部分 case 标为进行中以展示监控",
        ],
    )
    add_h(doc, "金标准（评测，非微调）", 2)
    add_bullets(
        doc,
        [
            "10–20 条：「给定制度片段 + 单据 → 期望下一节点/角色」人工标注，用于抽检准确率",
        ],
    )
    add_p(doc, "仓库目录建议：")
    add_code(
        doc,
        """data/policies/*.pdf
data/templates/*.yaml
data/cases/*.json
data/eval/gold.jsonl""",
    )

    add_h(doc, "六、核心模块与接口")
    add_p(doc, "建议工程结构：")
    add_bullets(
        doc,
        [
            "backend/app/ingest/ — PDF 解析与切块",
            "backend/app/rag/ — 检索",
            "backend/app/extract/ — 制度 → ProcessTemplate（JSON Schema 约束）",
            "backend/app/engine/ — next_step(case, template) / estimate_timeline(...)",
            "backend/app/bpmn/ — template → BPMN XML",
            "backend/app/report/ — LLM 报告生成",
            "backend/app/api/ — REST：/policies /templates /cases /portfolio /cases/{id}/timeline",
            "frontend/ — 项目群清单页 + 详情抽屉",
        ],
    )
    add_p(doc, "ProcessTemplate 最小字段示例：")
    add_code(
        doc,
        """id: expense_reimburse_v1
name: 费用报销
nodes:
  - id: submit
    name: 员工提交
    role: applicant
    sla_days: 0
  - id: dept_manager
    name: 部门经理审批
    role: dept_manager
    sla_days: 1
  - id: finance
    name: 财务审核
    role: finance
    sla_days: 2
gateways:
  - when: "amount >= 10000"
    then: cfo_approve
edges: [...]""",
    )
    add_p(doc, "推演输出示例（清单「下一步」列）：")
    add_p(
        doc,
        "2026-01-20，提交至财务审核，由财务专员审批（模板 SLA 2 个工作日）。",
        bold=True,
    )

    add_h(doc, "七、里程碑")
    add_p(
        doc,
        "原三周研发计划压缩为 W1/W2；另增 W3，专门将整套方法与约定封装为 Cursor Agent Skill（步骤细则见《技术栈》文档）。",
    )
    add_table(
        doc,
        ["周", "产出"],
        [
            [
                "W1",
                "数据目录与制度入库；模板 Schema；制度抽取流水线；推演引擎；合成清单 API；模拟提交接口",
            ],
            [
                "W2",
                "项目群清单前端；BPMN 渲染；分析报告；金标准抽检；演示脚本（导入制度 → 看清单 → 点开时间线/BPMN）",
            ],
            [
                "W3",
                "将方案与技术栈沉淀为可复用 Skill：梳理触发场景与硬约束 → 编写 SKILL.md（含 ProcessTemplate/目录/验收标准）→ 可选 scripts/reference → 安装到个人或项目 skills 目录 → 用真实对话验收「点名即按本方案搭建/扩展」",
            ],
        ],
    )
    add_p(
        doc,
        "说明：W3 不重复造业务功能，目标是让 AI 在后续项目中按同一套约定执行（选型 B+C、模板推演、PDF 学路径、不做 OA 挖掘等）。完整操作步骤写入《OA审批AI_技术栈》「封装为 Cursor Skill」一章。",
    )

    add_h(doc, "八、成功标准")
    add_bullets(
        doc,
        [
            "清单可展示 ≥8 条跨类型实例，列齐全。",
            "对金标准集，下一节点准确率 ≥85%（引擎命中模板，而非纯 LLM 自由生成）。",
            "报告/解释中制度引用可回溯到 PDF 片段。",
            "BPMN 能由模板稳定渲染（不依赖 LLM 直接画拓扑）。",
        ],
    )

    add_h(doc, "九、落地起步步骤")
    add_num(
        doc,
        [
            "初始化前后端骨架与 data/ 约定。",
            "先用 1 类流程（费用报销）跑通：PDF → 模板 → 合成 case → 清单 → 时间线/BPMN。",
            "再扩展采购/合同/请假，充实项目群清单。",
        ],
    )

    out = desk / "OA审批AI_项目方案.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
