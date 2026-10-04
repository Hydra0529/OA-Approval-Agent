from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Cm


desk = Path(r"C:\Users\Jhy13\Desktop")
EAST_ASIA = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia"


def _font(run, size=12, bold=False, italic=False, color=None):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(EAST_ASIA, "宋体")
    if color is not None:
        run.font.color.rgb = color


def add_title(doc, text):
    p = doc.add_heading(text, level=0)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0x1C, 0x24, 0x30)


def add_h(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_p(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _font(run, bold=bold)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    return p


def add_quote(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _font(run, size=11, italic=True, color=RGBColor(0x5C, 0x6B, 0x7A))
    return p


def add_bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(it, style="List Bullet")
        for run in p.runs:
            _font(run)


def add_num(doc, items):
    for it in items:
        p = doc.add_paragraph(it, style="List Number")
        for run in p.runs:
            _font(run)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                _font(run, size=11, bold=True)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    _font(run, size=11)
    doc.add_paragraph()


def setup_margins(doc):
    for s in doc.sections:
        s.top_margin = Cm(2.2)
        s.bottom_margin = Cm(2.2)
        s.left_margin = Cm(2.5)
        s.right_margin = Cm(2.5)


def build_plan():
    doc = Document()
    setup_margins(doc)
    add_title(doc, "公司 OA 审批 AI Agent —— 项目方案（通俗版）")
    add_quote(doc, "写给自己和答辩用的说明。尽量讲人话，少堆术语。")

    add_h(doc, "1. 这个项目是干什么的？")
    add_p(doc, "很多公司请假、报销、采购、签合同，都要在 OA 里走审批。")
    add_p(doc, "人经常搞不清：")
    add_bullets(doc, ["下一步该找谁批？", "大概还要多久？", "整条流程长什么样？"])
    add_p(
        doc,
        "这个项目做一个 AI 小助手：你提交（或模拟提交）一条审批后，它能告诉你下一步、预计时间，还能画出流程图，并给出简单分析。",
    )
    add_p(doc, "首期先做一张「项目群清单」——像一张总表，把好多条审批摆在一起看。")

    add_h(doc, "2. 它和视频里那个软件有什么关系？")
    add_p(
        doc,
        "参考视频里是「按规则推演」的流程软件：画好流程 → 按时间往后走 → 看节点和处理情况。",
    )
    add_p(doc, "我们不是照抄那套大型系统，而是学它的思路：")
    add_bullets(
        doc,
        [
            "先弄清楚流程怎么走；",
            "再根据当前进度推下一步和时间；",
            "最后用清单 + 流程图 + 报告展示出来。",
        ],
    )
    add_p(doc, "区别在于：视频主要靠人配规则；我们尽量让 AI 从公司规章制度 PDF 里读出规则。")

    add_h(doc, "3. 流程是怎么「学」来的？")
    add_p(doc, "重要约定：流程路径不靠真实 OA 历史数据学，而是从 PDF 规章制度学。")
    add_p(doc, "比如：")
    add_bullets(
        doc,
        [
            "《费用报销管理办法》",
            "《采购管理制度》",
            "《合同审批管理办法》",
            "《员工请假管理制度》",
        ],
    )
    add_p(doc, "AI / 程序读这些文件，整理成「流程模板」。例如报销大致是：")
    add_p(doc, "员工提交 → 部门经理 → 财务审核 →（钱多的话再找财务负责人）→ 出纳付款")
    add_p(
        doc,
        "清单里的例子数据可以自己编，用来演示；真正决定「下一步找谁」的，是制度抽出来的模板，不是瞎猜。",
    )

    add_h(doc, "4. 首期要做成什么样？（能演示就行）")
    add_p(doc, "打开网页，能看到一张表，每一行是一条审批，大致有：")
    add_table(
        doc,
        ["你能看到的", "含义"],
        [
            ["单据名称 / 类型", "这是报销还是采购等"],
            ["部门 / 申请人", "谁提的"],
            ["当前节点", "现在卡在哪一步"],
            ["下一步时间线", "例如：某日提交到财务审核，由财务专员审批"],
            ["预计时间 / 累计天数", "大概多久走完"],
            ["状态", "进行中 / 已完成"],
        ],
    )
    add_p(doc, "点开某一行，还能看：")
    add_bullets(
        doc,
        [
            "时间线：一步步什么时候该到谁；",
            "流程图：整条路长什么样；",
            "分析报告：流程说明、可能慢在哪、依据哪条制度。",
        ],
    )
    add_p(doc, "另外可以模拟提交一条新单子，立刻进清单——不用接公司真实 OA。")

    add_h(doc, "5. 整体怎么工作？（四步）")
    add_bullets(
        doc,
        [
            "规章制度 PDF",
            "读进去、切成小段（方便查找）",
            "整理成「流程模板」（节点、谁批、几天、金额条件）",
            "有一条审批进来 → 按模板算下一步和时间",
            "显示在项目群清单 + 流程图 + 报告里",
        ],
    )
    add_p(doc, "可以记成一句话：制度出规则，规则算进度，页面给人看。", bold=True)

    add_h(doc, "6. 为什么不全靠 AI「自由发挥」？")
    add_p(doc, "如果完全让大模型直接说「下一步是谁」，它有时会编。")
    add_p(doc, "所以做法是：")
    add_bullets(
        doc,
        [
            "算下一步、算日期：用模板和固定逻辑（可靠、能复查）；",
            "读制度、写说明、写报告：可以交给大模型（更会说话）。",
        ],
    )
    add_p(doc, "大模型不能擅自改审批状态，只帮忙理解和解释。")

    add_h(doc, "7. 我准备了哪些数据？")
    add_p(doc, "因为学校 / 演示用不了公司隐私数据，所以自己准备：")
    add_bullets(
        doc,
        [
            "制度 PDF：报销、采购、合同、请假等（演示版即可）；",
            "若干条假的审批例子：凑成项目群清单；",
            "一小份「标准答案」：用来检查「下一步算得对不对」。",
        ],
    )
    add_p(doc, "不涉及真实隐私，也不做模型训练 / 微调；需要云端模型时，用现成的开放接口就行。")

    add_h(doc, "8. 大概分三周做什么？")
    add_table(
        doc,
        ["周次", "目标（白话）"],
        [
            ["第 1 周", "搭好架子，制度进系统，清单能出假数据"],
            ["第 2 周", "能从制度抽流程、能推下一步、网页能提交和看表"],
            ["第 3 周", "能出流程图和报告，跑一遍正确率检查，准备演示"],
        ],
    )

    add_h(doc, "9. 做成什么样算过关？")
    add_bullets(
        doc,
        [
            "清单里至少有好几条不同类型的审批；",
            "「下一步」尽量算对（用标准答案测，希望大部分正确）；",
            "报告 / 说明能对应到制度内容；",
            "流程图是按模板稳定画出来的，不是模型随便画一笔。",
        ],
    )
    add_p(doc, "首期不做：")
    add_bullets(
        doc,
        [
            "对接公司真实 OA；",
            "从历史审批流水里挖流程；",
            "自己训练大模型；",
            "特别复杂的仿真播放（以后有空再加）。",
        ],
    )

    add_h(doc, "10. 一句话总结")
    add_p(
        doc,
        "做一个面向公司日常审批的 AI 演示系统：读懂制度 → 推演进度 → 用清单和流程图讲清楚。适合当课设 / 毕设前期 Demo，先讲清思路，再慢慢接真实系统。",
    )
    return doc


def build_stack():
    doc = Document()
    setup_margins(doc)
    add_title(doc, "公司 OA 审批 AI Agent —— 技术栈说明（通俗版）")
    add_quote(doc, "同样用大白话写：用了什么、各自干什么、为什么选它。")

    add_h(doc, "1. 先看一张总图")
    add_p(doc, "可以把整个项目想成几层：")
    add_table(
        doc,
        ["层", "干什么", "主要用什么"],
        [
            ["前台（网页）", "给人看清单、点详情、提交演示单", "React + TypeScript"],
            ["后台（服务）", "读 PDF、管模板、算下一步、出报告", "Python + FastAPI"],
            ["数据", "制度、模板、案例", "PDF / YAML / JSON 文件"],
            ["智能", "读制度、写报告（可选）", "云端大模型接口（DeepSeek / 通义等）"],
        ],
    )
    add_p(doc, "本地开发时大致是：网页 http://localhost:5173 ，接口 http://localhost:8000 。")

    add_h(doc, "2. 前台（你眼睛看到的部分）")
    add_h(doc, "React + TypeScript", 2)
    add_bullets(
        doc,
        [
            "React：做网页界面的常用框架，适合做表格、弹窗、详情页。",
            "TypeScript：带类型的 JavaScript，写的时候少犯低级错，课设也更规范一点。",
        ],
    )
    add_h(doc, "Vite", 2)
    add_p(doc, "启动前端开发服务器的工具，改代码刷新快，配置也简单。")
    add_h(doc, "bpmn-js", 2)
    add_bullets(
        doc,
        [
            "专门用来显示流程图的组件。",
            "后台按模板生成标准流程图文件，前台负责画出来给人看。",
            "首期只要「能看」，不要求在网页上拖拽改流程。",
        ],
    )
    add_h(doc, "界面大概长什么样", 2)
    add_bullets(
        doc,
        [
            "主页面：项目群清单表；",
            "右侧 / 抽屉：时间线、流程图、分析报告；",
            "按钮：刷新、模拟提交、跑一下正确率测试。",
        ],
    )

    add_h(doc, "3. 后台（真正算事的部分）")
    add_h(doc, "Python", 2)
    add_p(doc, "读文件、处理文本、调 AI 都方便，课设里很常见。")
    add_h(doc, "FastAPI", 2)
    add_bullets(
        doc,
        [
            "用 Python 写网页接口（API）的框架。",
            "浏览器 / 前端通过它问：「给我清单」「这条的下一步是什么」。",
            "自带文档页，方便自己调试。",
        ],
    )
    add_h(doc, "几个小模块（按功能分）", 2)
    add_table(
        doc,
        ["模块", "人话解释"],
        [
            ["读 PDF", "把制度文件里的字抠出来"],
            ["知识检索（简易版）", "按关键词找到相关制度段落（演示够用）"],
            ["流程抽取", "从制度整理成「模板」：几步、谁批、几天、金额条件"],
            ["推演引擎", "按模板算下一步和日期（可靠的那部分）"],
            ["流程图生成", "把模板变成流程图文件"],
            ["报告生成", "汇总情况，写成一段分析（有 AI 用 AI，没有就用固定模板写）"],
        ],
    )

    add_h(doc, "4. 数据和「流程模板」存哪？")
    add_p(doc, "首期先不搞复杂数据库，用文件夹就够演示：")
    add_table(
        doc,
        ["目录", "放什么"],
        [
            ["data/policies/", "规章制度 PDF"],
            ["data/templates/", "流程模板（YAML，人也能打开看）"],
            ["data/cases/", "模拟的审批例子（JSON）"],
            ["data/eval/", "检查「下一步对不对」的小测验题"],
        ],
    )
    add_p(doc, "以后要正经上线，可以再换成正式数据库；课设阶段文件更直观。")

    add_h(doc, "5. AI / 大模型怎么用？")
    add_h(doc, "用什么", 2)
    add_bullets(
        doc,
        [
            "云端开放接口即可，例如 DeepSeek、通义千问等。",
            "不做微调、不自己训练模型（省钱、省事，也符合「不碰隐私训练」）。",
        ],
    )
    add_h(doc, "用在哪", 2)
    add_bullets(
        doc,
        [
            "帮着从制度里整理流程草案；",
            "写分析报告、解释「为什么走这条」；",
            "制度问答（可选）。",
        ],
    )
    add_h(doc, "不用在哪", 2)
    add_bullets(
        doc,
        [
            "不算最终审批路由（下一步找谁，以模板引擎为准）；",
            "不直接改审批通过 / 驳回。",
        ],
    )
    add_h(doc, "没有 API Key 怎么办？", 2)
    add_bullets(
        doc,
        [
            "系统也能跑：用写好的演示模板 + 规则降级；",
            "有 Key 时再打开「更聪明」的抽取和报告。",
        ],
    )

    add_h(doc, "6. 为什么选这套，而不是别的？")
    add_table(
        doc,
        ["选择", "简单理由"],
        [
            ["Python 后台", "处理文档和 AI 方便，资料多"],
            ["React 前台", "做管理表格类页面成熟"],
            ["文件存数据", "演示快，不依赖装数据库"],
            ["模板引擎 + AI", "既稳又会说话，适合审批这种不能乱猜的场景"],
            ["云端大模型", "不用自己买显卡训练"],
        ],
    )
    add_p(doc, "一句话：学生课设能跑通、逻辑说得清、以后还能扩展。", bold=True)

    add_h(doc, "7. 我电脑上要准备什么？")
    add_p(doc, "大致需要：")
    add_bullets(
        doc,
        [
            "Python 3.11+（跑后台）；",
            "Node.js + npm（跑前台）；",
            "（可选）大模型 API Key，写在后台的 .env 里。",
        ],
    )
    add_p(doc, "常用步骤：后台安装依赖并启动接口服务；前台 npm install 后 npm run dev 打开网页。")

    add_h(doc, "8. 以后可以怎么升级？（了解即可）")
    add_bullets(
        doc,
        [
            "接真实 OA 的消息通知（有人提交就自动进清单）；",
            "换成更专业的向量数据库做制度检索；",
            "有大量历史数据后，再做流程分析加强（不是首期重点）。",
        ],
    )

    add_h(doc, "9. 封装为 Cursor Skill（对应方案里程碑 W3）")
    add_p(
        doc,
        "目标：把本项目的选型、目录约定、模板字段、推演硬约束、验收标准写成 AI 可执行说明书，以后点名该 Skill，就能按同一套方法搭建或扩展，而不是把整仓代码贴进 Skill。",
    )

    add_h(doc, "9.1 Skill 是什么、装在哪", 2)
    add_bullets(
        doc,
        [
            "Skill = 一个文件夹，核心是 SKILL.md（可附带 reference.md、examples.md、scripts/）。",
            "个人技能（本机所有项目可用）：~/.cursor/skills/<skill-name>/（Windows 一般在 C:\\Users\\<你>\\.cursor\\skills\\）。",
            "项目技能（跟着仓库走）：<仓库>/.cursor/skills/<skill-name>/。",
            "不要写到 ~/.cursor/skills-cursor/（那是 Cursor 内置技能目录）。",
        ],
    )

    add_h(doc, "9.2 封装前先定清楚的 5 件事", 2)
    add_num(
        doc,
        [
            "用途：例如「按本方案搭建/扩展 OA 审批 AI Agent（制度 PDF → 模板 → 推演 → 项目群清单）」。",
            "触发词：用户提到 OA 审批 Agent、制度抽流程、项目群清单、BPMN 报告、SLA 推演等时启用。",
            "存放位置：个人 or 项目级（实习汇报建议先项目级，方便随仓库提交）。",
            "必须写进 Skill 的领域知识：方案 B+C、路径只从 PDF 学、LLM 不算路由、目录与 ProcessTemplate 字段、成功标准。",
            "不要塞进 Skill 的：整段前后端源码；改为「打开本仓库某路径 / 运行某命令」。",
        ],
    )

    add_h(doc, "9.3 建议目录结构", 2)
    for line in [
        "oa-approval-agent/",
        "├── SKILL.md                 # 必填：何时用、硬约束、落地步骤",
        "├── reference.md            # 选填：模板字段、API 一览、里程碑",
        "├── examples.md             # 选填：用户怎么说、Agent 应怎么做",
        "└── scripts/                # 选填：如 seed_demo、eval 检查脚本入口说明",
    ]:
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(10)
        run._element.rPr.rFonts.set(EAST_ASIA, "宋体")
        p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph()

    add_h(doc, "9.4 逐步操作清单（按顺序做）", 2)
    add_num(
        doc,
        [
            "从《项目方案》抽出「不可违反的约定」清单（路径来源、LLM 边界、首期不做项）。",
            "从本文技术栈抽出「默认选型表」与模块职责（前端 React、后端 FastAPI、推演引擎、BPMN 等）。",
            "新建 skill 目录（项目级优先：.cursor/skills/oa-approval-agent/）。",
            "编写 SKILL.md 的 YAML 头：name（小写+连字符，如 oa-approval-agent）；description 用第三人称，写清 WHAT + WHEN，并带上触发关键词。",
            "默认可加 disable-model-invocation: true——只有用户点名或明确相关时再加载，避免乱触发；若希望自动匹配场景，再去掉该项。",
            "编写 SKILL.md 正文：硬约束 → 推荐架构（B 内核 + C 交互）→ 目录约定 → ProcessTemplate 最小字段 → 落地步骤（对齐 W1/W2）→ 验收标准 →「不要做」列表。",
            "正文保持短：只写 Agent 不容易猜到的项目约定；通用编程知识少写。",
            "（可选）写 reference.md：YAML 模板样例、REST 路径列表、金标准评测怎么跑。",
            "（可选）写 examples.md：2～3 条对话示例（「从零搭 Demo」「只加一类请假流程」「把制度 PDF 抽成模板」）。",
            "（可选）在 scripts/ 放薄封装或说明：指向仓库已有 seed / eval 命令，而不是复制大段业务代码。",
            "保存后重启 Cursor 或新开 Agent 对话，确认技能出现在可用 Skills 列表。",
            "验收对话 1：点名 Skill，要求「按本方案搭一个空骨架」——检查是否遵守目录与硬约束。",
            "验收对话 2：要求「只根据制度增加采购分支」——检查是否走模板+引擎，而不是让 LLM 直接猜下一步。",
            "验收对话 3：故意说「用 OA 日志挖流程」——检查 Skill 是否明确拒绝并指向首期不做项。",
            "根据验收结果改短 description / 补硬约束，避免 Skill 又长又空。",
            "若需给组员用：把 .cursor/skills/... 提交进 Git；若仅自己用：复制到 ~/.cursor/skills/。",
        ],
    )

    add_h(doc, "9.5 SKILL.md 建议必备章节（提纲）", 2)
    add_bullets(
        doc,
        [
            "When to use：触发场景。",
            "Hard constraints：PDF 学路径；引擎算下一步；不微调；不接真 OA（首期）。",
            "Default stack：与本文第 1～3 节一致的选型表。",
            "Repo layout：data/policies|templates|cases|eval，backend/，frontend/。",
            "ProcessTemplate schema：nodes / edges / gateways / sla_days / role。",
            "Implementation order：先报销一类跑通，再扩类型；对齐方案 W1→W2。",
            "Definition of done：清单≥8 条、金标准≥85%、BPMN 由模板生成。",
            "Out of scope：流程挖掘优先、Agent 静默改签等。",
        ],
    )

    add_h(doc, "9.6 description 写法示例（可直接改）", 2)
    add_p(
        doc,
        "Guides building an OA approval AI Agent that learns process paths from policy PDFs, uses deterministic template engines for next-step/SLA timelines, and delivers a portfolio list with BPMN and reports. Use when the user mentions OA审批, 制度抽流程, 项目群清单, BPMN 审批流程, or packaging this internship agent architecture.",
    )

    add_h(doc, "9.7 常见坑", 2)
    add_bullets(
        doc,
        [
            "把整仓代码贴进 SKILL.md → 又长又难维护；应引用路径与命令。",
            "description 太虚（「帮助做审批」）→ Agent 不知道何时调用。",
            "漏写硬约束 → Agent 容易改成「纯 LLM 猜路由」或「上流程挖掘」。",
            "装错目录（skills-cursor）→ 技能不会被当作用户 Skill 加载。",
        ],
    )

    add_h(doc, "10. 一句话记住技术栈")
    add_p(
        doc,
        "React 做页面，Python / FastAPI 做脑子，PDF 出规则，模板算进度，大模型帮忙读和写；稳定后把约定收成 Skill，给以后的 AI 按同一套规矩干活。",
        bold=True,
    )
    return doc


def main():
    path2 = desk / "OA审批AI_技术栈.docx"
    build_stack().save(path2)
    print(path2)


if __name__ == "__main__":
    main()
